import io

import docx
import pytest
from PIL import Image, ImageDraw, ImageFont

from stylelens.eval import samples
from stylelens.ingestion.ocr import ocr_available
from stylelens.ingestion.parsers import get_parser
from stylelens.ingestion.parsers.base import assign_heading_paths, noop_progress


@pytest.fixture(scope="module")
def northwind():
    return get_parser("pdf").parse(samples.northwind_guide(), "northwind.pdf", noop_progress)


def test_pdf_structure(northwind):
    types = {e.type for e in northwind.elements}
    assert {"heading", "paragraph", "list_item", "table", "image", "swatch"} <= types
    assert northwind.title == "Northwind Coffee Brand Guidelines"
    assert len(northwind.pages) >= 5
    assert all(p.render and p.render.startswith(b"\x89PNG") for p in northwind.pages)
    # every element is positioned on a real page
    for el in northwind.elements:
        assert 1 <= el.page <= len(northwind.pages)
        assert el.bbox is not None


def test_pdf_heading_paths(northwind):
    paths = assign_heading_paths(northwind.elements)
    by_text = {e.text: p for e, p in zip(northwind.elements, paths)}
    clear = next(p for t, p in by_text.items() if t.startswith("Keep clear space around the logo"))
    assert clear[-2:] == ["5. Logo", "Clear Space"]
    oxford = next(p for t, p in by_text.items() if "serial (Oxford) comma" in t)
    assert "7. Writing Style" in oxford and oxford[-1] == "Grammar and mechanics"


def test_pdf_swatches_and_labels(northwind):
    fills = [e.meta["fill"] for e in northwind.elements if e.type == "swatch"]
    assert fills[:4] == ["#1B4F9C", "#3B2A20", "#F2E3C6", "#D9622B"]
    # grid labels are not mistaken for headings
    assert not any(e.type == "heading" and e.text == "Harbor Blue" for e in northwind.elements)


def test_pdf_lists_and_tables(northwind):
    items = [e.text for e in northwind.elements if e.type == "list_item"]
    assert any(t.startswith("Use the serial (Oxford) comma") for t in items)
    scale = next(e for e in northwind.elements if e.type == "table" and "Type" in e.text and "H1" in e.text)
    assert "48px / 56px" in scale.text
    assert scale.meta["rows"][1][0] == "H1"


def test_pdf_font_metadata(northwind):
    heading = next(e for e in northwind.elements if e.type == "heading" and e.text == "4. Typography")
    assert heading.meta["fonts"][0]["bold"] is True
    assert heading.meta["fonts"][0]["size"] > northwind.meta["body_font_size"]
    assert "Helvetica" in northwind.meta["embedded_fonts"]


def test_docx_parser():
    d = docx.Document()
    d.add_heading("Acme Style Guide", 0)
    d.add_heading("Colors", 1)
    d.add_paragraph("Primary color is #FF0000.")
    d.add_paragraph("Never use green.", style="List Bullet")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Name", "Hex"
    t.cell(1, 0).text, t.cell(1, 1).text = "Red", "#FF0000"
    buf = io.BytesIO()
    d.save(buf)
    doc = get_parser("docx").parse(buf.getvalue(), "acme.docx", noop_progress)
    types = [e.type for e in doc.elements]
    assert types == ["heading", "heading", "paragraph", "list_item", "table"]
    paths = assign_heading_paths(doc.elements)
    assert paths[2] == ["Acme Style Guide", "Colors"]
    assert "| Red | #FF0000 |" in doc.elements[-1].text


def test_markdown_parser():
    md = b"# Guide\n\nIntro text\nwraps here.\n\n## Voice\n\n- Be kind\n- Be brief\n\n| a | b |\n|---|---|\n| 1 | 2 |\n"
    doc = get_parser("markdown").parse(md, "g.md", noop_progress)
    assert [e.type for e in doc.elements] == ["heading", "paragraph", "heading", "list_item", "list_item", "table"]
    assert doc.elements[1].text == "Intro text wraps here."
    assert doc.elements[1].meta == {"line_start": 3, "line_end": 4}


def test_text_parser_headings():
    doc = get_parser("text").parse(b"COLOR RULES\n\nUse blue.\n", "g.txt", noop_progress)
    assert doc.elements[0].type == "heading" and doc.elements[0].text == "Color Rules"


@pytest.mark.skipif(not ocr_available(), reason="tesseract not installed")
def test_image_parser_ocr():
    img = Image.new("RGB", (900, 200), "white")
    ImageDraw.Draw(img).text((20, 60), "Never stretch the logo", fill="black", font=ImageFont.load_default(size=56))
    buf = io.BytesIO()
    img.save(buf, "PNG")
    doc = get_parser("image").parse(buf.getvalue(), "rule.png", noop_progress)
    assert doc.elements[0].type == "image"
    ocr = " ".join(e.text for e in doc.elements[1:]).lower()
    assert "stretch" in ocr and "logo" in ocr


def test_svg_parser():
    svg = b'<svg xmlns="http://www.w3.org/2000/svg" width="200" height="100"><rect width="80" height="80" fill="#1B4F9C"/><text x="90" y="50" font-size="14">Harbor</text></svg>'
    doc = get_parser("svg").parse(svg, "logo.svg", noop_progress)
    assert any(e.type == "swatch" and e.meta["fill"] == "#1B4F9C" for e in doc.elements)


def test_pptx_parser():
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Color palette"
    body = s.placeholders[1].text_frame
    body.text = "Primary is Harbor Blue #1B4F9C"
    p = body.add_paragraph()
    p.text = "Never use pure black"
    p.level = 1
    chip = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(5), Inches(1), Inches(1))
    chip.fill.solid()
    chip.fill.fore_color.rgb = RGBColor(0x1B, 0x4F, 0x9C)
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Type scale"
    tbl = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    tbl.cell(0, 0).text, tbl.cell(0, 1).text = "Style", "Typeface"
    tbl.cell(1, 0).text, tbl.cell(1, 1).text = "H1", "Fraunces"
    buf = io.BytesIO()
    prs.save(buf)
    doc = get_parser("pptx").parse(buf.getvalue(), "deck.pptx", noop_progress)
    assert len(doc.pages) == 2
    assert doc.elements[0].type == "heading" and doc.elements[0].text == "Color palette"
    assert any(e.type == "swatch" and e.meta["fill"] == "#1B4F9C" for e in doc.elements)
    assert any(e.type == "table" and "Fraunces" in e.text and e.page == 2 for e in doc.elements)
    assert any("pure black" in e.text for e in doc.elements)
    assert all(e.bbox for e in doc.elements if e.type != "paragraph" or not e.meta.get("speaker_notes"))


def test_html_parser():
    html = b"""<html><head><title>Acme Web Guide</title>
    <style>:root{--brand:#0055FF} body{font-family: "Inter", sans-serif; color:#222}</style></head>
    <body><h1>Voice</h1><p>We are <b>clear</b> and kind.</p><ul><li>Never use jargon.</li><li>Use "you".</li></ul>
    <h2>Buttons</h2><table><tr><th>State</th><th>Color</th></tr><tr><td>Default</td><td>#0055FF</td></tr></table>
    <img src="logo.png" alt="Acme logo on white"><script>var x = "#FF0000";</script></body></html>"""
    doc = get_parser("html").parse(html, "guide.html", noop_progress)
    assert doc.title == "Acme Web Guide"
    types = [e.type for e in doc.elements]
    assert types[:5] == ["heading", "paragraph", "list_item", "list_item", "heading"]
    assert doc.elements[1].text == "We are clear and kind."
    assert any(e.type == "table" and "#0055FF" in e.text for e in doc.elements)
    css = [e.text for e in doc.elements if e.meta.get("css")]
    assert any("Typeface: Inter" in c for c in css) and any("#0055FF" in c for c in css)
    assert not any("#FF0000" in e.text for e in doc.elements)  # scripts ignored
    paths = assign_heading_paths(doc.elements)
    assert paths[-1] == ["Stylesheet typography and colors"]
