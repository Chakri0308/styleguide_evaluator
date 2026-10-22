"""PPTX parser (python-pptx): one page per slide; titles become headings, body text keeps bullet levels,
tables become Markdown, pictures become images, and solid-filled shapes become colour swatches."""

import io

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.util import Emu

from stylelens.ingestion.parsers.base import (
    HEADING,
    IMAGE,
    LIST_ITEM,
    PARAGRAPH,
    SWATCH,
    TABLE,
    ParsedDocument,
    ParsedElement,
    ParsedPage,
    ProgressFn,
)
from stylelens.ingestion.parsers.pdf import _table_markdown

EMU_PER_PT = 12700
TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}


def _bbox(shape) -> tuple[float, float, float, float] | None:
    if shape.left is None or shape.top is None:
        return None
    x0, y0 = Emu(shape.left) / EMU_PER_PT, Emu(shape.top) / EMU_PER_PT
    return (x0, y0, x0 + Emu(shape.width or 0) / EMU_PER_PT, y0 + Emu(shape.height or 0) / EMU_PER_PT)


def _fonts(paragraph) -> list[dict]:
    out, seen = [], set()
    for run in paragraph.runs:
        f = run.font
        color = None
        try:
            if f.color is not None and f.color.type is not None and f.color.rgb is not None:
                color = f"#{f.color.rgb}"
        except (AttributeError, ValueError):
            pass
        key = (f.name, f.size.pt if f.size else None, bool(f.bold), color)
        if key not in seen and run.text.strip():
            seen.add(key)
            out.append({"font": key[0], "size": key[1], "bold": key[2], "color": color, "chars": len(run.text)})
    return out


def _fill_hex(shape) -> str | None:
    try:
        fill = shape.fill
        if fill.type == 1 and fill.fore_color and fill.fore_color.type is not None:  # MSO_FILL.SOLID
            return f"#{fill.fore_color.rgb}"
    except (AttributeError, ValueError, TypeError):
        return None
    return None


class PptxParser:
    name = "python-pptx"

    def parse(self, data: bytes, filename: str, progress: ProgressFn) -> ParsedDocument:
        prs = Presentation(io.BytesIO(data))
        width = Emu(prs.slide_width or 9144000) / EMU_PER_PT
        height = Emu(prs.slide_height or 6858000) / EMU_PER_PT
        pages: list[ParsedPage] = []
        elements: list[ParsedElement] = []
        n = len(prs.slides)
        for i, slide in enumerate(prs.slides):
            num = i + 1
            pages.append(ParsedPage(page=num, width=width, height=height))
            shapes = sorted(slide.shapes, key=lambda s: ((s.top or 0), (s.left or 0)))
            title_shape = slide.shapes.title
            if title_shape is not None and title_shape.has_text_frame and title_shape.text_frame.text.strip():
                elements.append(ParsedElement(HEADING, title_shape.text_frame.text.strip(), num, _bbox(title_shape),
                                              level=1, meta={"fonts": _fonts(title_shape.text_frame.paragraphs[0])}))
            for shape in shapes:
                if title_shape is not None and shape.shape_id == title_shape.shape_id:
                    continue
                box = _bbox(shape)
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    try:
                        img = shape.image
                        elements.append(ParsedElement(IMAGE, "", num, box, image=img.blob, image_ext=img.ext,
                                                      meta={"width_px": img.size[0], "height_px": img.size[1]}))
                    except (AttributeError, ValueError):
                        pass
                    continue
                if getattr(shape, "has_table", False):
                    rows = [[c.text for c in r.cells] for r in shape.table.rows]
                    md = _table_markdown(rows)
                    if md:
                        elements.append(ParsedElement(TABLE, md, num, box, meta={"rows": rows}))
                    continue
                fill = _fill_hex(shape)
                if fill and box and (box[2] - box[0]) >= 10 and (box[3] - box[1]) >= 10:
                    elements.append(ParsedElement(SWATCH, "", num, box, meta={"fill": fill}))
                if not getattr(shape, "has_text_frame", False) or not shape.text_frame.text.strip():
                    continue
                for para in shape.text_frame.paragraphs:
                    text = "".join(r.text for r in para.runs).strip() or para.text.strip()
                    if not text:
                        continue
                    kind = LIST_ITEM if para.level > 0 or len(shape.text_frame.paragraphs) > 2 else PARAGRAPH
                    elements.append(ParsedElement(kind, text, num, box, meta={"fonts": _fonts(para), "level": para.level}))
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                elements.append(ParsedElement(PARAGRAPH, slide.notes_slide.notes_text_frame.text.strip(), num, None,
                                              meta={"speaker_notes": True}))
            progress(num / max(n, 1), f"Slide {num} of {n}")
        title = prs.core_properties.title or next((e.text for e in elements if e.type == HEADING), None)
        return ParsedDocument(parser=self.name, title=title, pages=pages, elements=elements,
                              meta={"slides": n, "paged": False})
