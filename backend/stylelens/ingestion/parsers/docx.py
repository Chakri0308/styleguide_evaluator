"""DOCX parser (python-docx). Pages are estimated from rendered/explicit page breaks; no bboxes."""

import io
import re

import docx
from docx.oxml.ns import qn

from stylelens.ingestion.parsers.base import (
    HEADING,
    IMAGE,
    LIST_ITEM,
    PARAGRAPH,
    TABLE,
    ParsedDocument,
    ParsedElement,
    ParsedPage,
    ProgressFn,
)
from stylelens.ingestion.parsers.pdf import _table_markdown

LETTER = (612.0, 792.0)


def _run_fonts(paragraph) -> list[dict]:
    out, seen = [], set()
    style_font = paragraph.style.font if paragraph.style is not None else None
    for run in paragraph.runs:
        if not run.text.strip():
            continue
        f = run.font
        name = f.name or (style_font.name if style_font else None)
        size = f.size or (style_font.size if style_font else None)
        bold = bool(f.bold if f.bold is not None else (style_font.bold if style_font else False))
        color = None
        try:
            if f.color is not None and f.color.type is not None and f.color.rgb is not None:
                color = f"#{f.color.rgb}"
        except (AttributeError, ValueError):
            color = None
        key = (name, size.pt if size else None, bold, color)
        if key in seen:
            continue
        seen.add(key)
        out.append({"font": name, "size": key[1], "bold": bold, "color": color, "chars": len(run.text)})
    return out


class DocxParser:
    name = "python-docx"

    def parse(self, data: bytes, filename: str, progress: ProgressFn) -> ParsedDocument:
        d = docx.Document(io.BytesIO(data))
        elements: list[ParsedElement] = []
        page = 1
        body = d.element.body
        children = list(body.iterchildren())
        para_by_el = {p._p: p for p in d.paragraphs}
        table_by_el = {t._tbl: t for t in d.tables}

        for i, child in enumerate(children):
            if child.tag == qn("w:p"):
                p = para_by_el.get(child)
                if p is None:
                    continue
                breaks = len(child.findall(".//" + qn("w:lastRenderedPageBreak")))
                breaks += sum(1 for br in child.findall(".//" + qn("w:br")) if br.get(qn("w:type")) == "page")
                page += breaks
                # Inline images
                for blip in child.findall(".//" + qn("a:blip")):
                    rid = blip.get(qn("r:embed"))
                    part = d.part.related_parts.get(rid) if rid else None
                    if part is not None and getattr(part, "blob", None):
                        ext = part.partname.ext if hasattr(part.partname, "ext") else "png"
                        elements.append(ParsedElement(IMAGE, "", page, None, image=part.blob, image_ext=ext))
                text = p.text.strip()
                if not text:
                    continue
                style = (p.style.name if p.style is not None else "") or ""
                fonts = _run_fonts(p)
                m = re.match(r"Heading (\d)", style)
                if style == "Title":
                    elements.append(ParsedElement(HEADING, text, page, None, level=0, meta={"fonts": fonts, "style": style}))
                elif m:
                    elements.append(ParsedElement(HEADING, text, page, None, level=int(m.group(1)),
                                                  meta={"fonts": fonts, "style": style}))
                elif "List" in style or child.find(".//" + qn("w:numPr")) is not None:
                    elements.append(ParsedElement(LIST_ITEM, text, page, None, meta={"fonts": fonts, "style": style}))
                else:
                    elements.append(ParsedElement(PARAGRAPH, text, page, None, meta={"fonts": fonts, "style": style}))
            elif child.tag == qn("w:tbl"):
                t = table_by_el.get(child)
                if t is None:
                    continue
                rows = [[c.text for c in r.cells] for r in t.rows]
                md = _table_markdown(rows)
                if md:
                    elements.append(ParsedElement(TABLE, md, page, None, meta={"rows": rows}))
            progress((i + 1) / max(len(children), 1), f"Block {i + 1} of {len(children)}")

        title = d.core_properties.title or next((e.text for e in elements if e.type == HEADING), None)
        pages = [ParsedPage(page=n, width=LETTER[0], height=LETTER[1]) for n in range(1, page + 1)]
        styles = sorted({s.name for s in d.styles if getattr(s, "font", None) is not None and s.font.name})
        return ParsedDocument(parser=self.name, title=title, pages=pages, elements=elements,
                              meta={"styles_with_fonts": styles, "paged": False})
